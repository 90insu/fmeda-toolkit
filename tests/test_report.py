"""Report generation.

The workbook's metric cells are Excel formulas, not values computed in Python.
That is only safe if something checks the two agree — which is what
`test_excel_formulas_match_python` does, by recalculating the file with
LibreOffice and reading the cached results back.
"""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fmeda import policy as P
from fmeda.metrics import compute
from fmeda.models import load_analysis
from fmeda.report import write_report

EXAMPLE = Path(__file__).parent.parent / "examples" / "bjb" / "fmeda.yaml"
UPLOADS = [{"category": "schematic", "name": "bjb.net", "digest": "0" * 64}]
DOCS = set(P.DocCategory) - {P.DocCategory.SAFETY_MANUAL}


@pytest.fixture(scope="module")
def workbook_path(tmp_path_factory) -> Path:
    analysis, parts = load_analysis(EXAMPLE)
    result = P.apply(analysis, parts, P.Level.CONSERVATIVE, DOCS)
    goals = compute(result.analysis, parts)
    out = tmp_path_factory.mktemp("report") / "report.xlsx"
    return write_report(out, analysis, parts, result, goals, UPLOADS, DOCS,
                        "SN 29500", "Underhood, 40 C")


def test_every_expected_sheet_exists(workbook_path):
    names = load_workbook(workbook_path).sheetnames
    assert names[0] == "Summary", "Summary must open first"
    for expected in ("SG1", "SG2", "Mechanisms", "Policy", "Review queue",
                     "Assumptions", "Inputs"):
        assert expected in names


def test_goal_sheet_has_a_rationale_on_every_row(workbook_path):
    ws = load_workbook(workbook_path)["SG1"]
    rows = 0
    for row in ws.iter_rows(min_row=6, max_col=16):
        if row[0].value in (None, "TOTAL"):
            break
        assert row[15].value, f"row {row[0].row} has no rationale"
        rows += 1
    assert rows > 20


def test_mechanism_sheet_shows_stated_and_applied_coverage(workbook_path):
    """An assessor has to be able to see what the policy changed, not just the
    number it left behind."""
    ws = load_workbook(workbook_path)["Mechanisms"]
    headers = [c.value for c in ws[3]]
    assert "DC_SPF stated" in headers and "DC_SPF applied" in headers
    stated = headers.index("DC_SPF stated") + 1
    applied = headers.index("DC_SPF applied") + 1
    assert any(ws.cell(row=r, column=stated).value != ws.cell(row=r, column=applied).value
               for r in range(4, ws.max_row + 1))


def test_inputs_sheet_records_what_was_missing(workbook_path):
    text = " ".join(
        str(c.value) for row in load_workbook(workbook_path)["Inputs"].iter_rows()
        for c in row if c.value
    )
    assert "safety_manual" in text
    assert "SN 29500" in text


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not available")
def test_excel_formulas_match_python(workbook_path, tmp_path):
    out_dir = tmp_path / "recalculated"
    out_dir.mkdir()
    subprocess.run(
        ["soffice", "--headless", "-env:UserInstallation=file:///tmp/lo_fmeda_test",
         "--convert-to", "xlsx", "--outdir", str(out_dir), str(workbook_path)],
        check=True, capture_output=True, timeout=240,
    )
    recalculated = out_dir / workbook_path.name
    assert recalculated.exists(), "LibreOffice produced no output"

    ws = load_workbook(recalculated, data_only=True)["Summary"]
    analysis, parts = load_analysis(EXAMPLE)
    result = P.apply(analysis, parts, P.Level.CONSERVATIVE, DOCS)
    goals = compute(result.analysis, parts)

    for i, goal in enumerate(goals.values()):
        row = 5 + i
        assert ws.cell(row=row, column=3).value == pytest.approx(goal.spfm, abs=1e-9)
        assert ws.cell(row=row, column=5).value == pytest.approx(goal.lfm, abs=1e-9)
        assert ws.cell(row=row, column=7).value == pytest.approx(goal.pmhf_fit, abs=1e-6)
