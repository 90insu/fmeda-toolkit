"""Parity check: the browser build must agree with the Python core.

Two implementations of the same metric is a real risk, so this is the guard.
It drives the page in a headless browser, reads the numbers it renders,
compares them against fmeda.metrics, downloads the workbook the page writes,
and opens it with openpyxl.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))
sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from validate_ooxml import validate as validate_ooxml

from fmeda import policy as P
from fmeda.metrics import compute
from fmeda.models import load_analysis

PAGE = Path(__file__).parent.parent / "docs" / "index.html"
EXAMPLE = Path(__file__).parent.parent / "examples" / "bjb" / "fmeda.yaml"
OUT = Path(__file__).parent.parent / ".verify-out"
OUT.mkdir(parents=True, exist_ok=True)

analysis, parts = load_analysis(EXAMPLE)
# The page starts with nothing uploaded, so no document is present.
NO_DOCS: set = set()

failures = []

with sync_playwright() as pw:
    browser = pw.chromium.launch(executable_path="/opt/pw-browsers/chromium")
    ctx = browser.new_context(accept_downloads=True, viewport={"width": 1400, "height": 1200})
    page = ctx.new_page()
    errors = []
    page.on("pageerror", lambda e: errors.append(str(e)))
    # Google Fonts is blocked in this sandbox and will be unreachable whenever
    # the page is opened offline, which is the normal case. The stylesheet
    # failing is expected; the page declares real fallback stacks for it.
    page.on("console", lambda m: errors.append(m.text)
            if m.type == "error" and "ERR_TUNNEL" not in m.text
            and "net::" not in m.text else None)
    page.goto(PAGE.as_uri(), wait_until="load")
    page.wait_for_timeout(2500)

    if errors:
        failures.append(f"javascript errors: {errors[:4]}")

    for level in ("nominal", "conservative", "worst_case"):
        page.select_option("#level", level)
        page.wait_for_timeout(500)
        shown = page.evaluate(
            "() => Object.fromEntries(Object.values(CURRENT.goals).map("
            "g => [g.id, {spfm: g.spfm, lfm: g.lfm, pmhf: g.pmhf}]))"
        )
        expected = compute(P.apply(analysis, parts, P.Level(level), NO_DOCS).analysis, parts)
        for gid, exp in expected.items():
            got = shown[gid]
            for key, want in (("spfm", exp.spfm), ("lfm", exp.lfm), ("pmhf", exp.pmhf_fit)):
                if abs(got[key] - want) > 1e-9:
                    failures.append(
                        f"{level} {gid} {key}: browser {got[key]!r} != python {want!r}"
                    )
        print(f"  {level:<13} " + "  ".join(
            f"{g}: SPFM {v['spfm']*100:6.2f}% LFM {v['lfm']*100:6.2f}% PMHF {v['pmhf']:6.2f}"
            for g, v in shown.items()))

    # Back to the default the page ships with, then take the workbook.
    page.select_option("#level", "conservative")
    page.wait_for_timeout(400)
    with page.expect_download(timeout=30000) as dl:
        page.click("#xlsx")
    download = dl.value
    xlsx = OUT / "browser-report.xlsx"
    download.save_as(xlsx)
    print(f"\n  downloaded {download.suggested_filename} ({xlsx.stat().st_size} bytes)")

    page.screenshot(path=str(OUT / "bench.png"), full_page=False)
    browser.close()

# Excel enforces OOXML element ordering that openpyxl and LibreOffice ignore,
# so structural validation has to happen before the lenient readers get a look.
structural = validate_ooxml(xlsx)
if structural:
    failures.extend(f"ooxml: {p}" for p in structural)
    print("  ooxml structure: INVALID")
else:
    print("  ooxml structure: valid")

wb = load_workbook(xlsx)
print("  sheets:", wb.sheetnames)
for expected_sheet in ("Summary", "SG1", "SG2", "Mechanisms", "Policy",
                       "Review queue", "Assumptions", "Inputs"):
    if expected_sheet not in wb.sheetnames:
        failures.append(f"workbook missing sheet {expected_sheet!r}")

sg1 = wb["SG1"]
rationales = [sg1.cell(row=r, column=16).value for r in range(6, 40)]
filled = [x for x in rationales if x]
print(f"  SG1 rationale cells filled: {len(filled)}")
if len(filled) < 20:
    failures.append(f"only {len(filled)} rationale cells")

formulas = [c.value for row in wb["Summary"].iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")]
print(f"  Summary formulas: {len(formulas)} e.g. {formulas[0] if formulas else 'NONE'}")
if len(formulas) < 6:
    failures.append("summary metric formulas missing")

print()
if failures:
    print("FAILURES:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("parity OK - browser matches the Python core, workbook opens clean")
